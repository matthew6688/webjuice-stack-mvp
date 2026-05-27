#!/usr/bin/env python3
"""
Helper: convert NSW Fair Trading XLSX (multi-sheet) → single CSV.
Used by pl:license-csv-sync.

Usage:
  python3 _xlsx-to-csv.py <input.xlsx> <output.csv>
"""
import openpyxl, csv, sys

if len(sys.argv) != 3:
    print("Usage: _xlsx-to-csv.py <in.xlsx> <out.csv>", file=sys.stderr)
    sys.exit(2)

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True)

# Find sheets that look like data sheets (with header row matching pattern)
HEADER_HINTS = ['Licence Number', 'Licensee']
data_sheets = []
for sn in wb.sheetnames:
    ws = wb[sn]
    rows_iter = ws.iter_rows(values_only=True)
    # Skip first row · check 2nd row for header pattern
    first = next(rows_iter, None)
    second = next(rows_iter, None)
    if second and any(h in (second or []) for h in HEADER_HINTS):
        data_sheets.append(sn)

if not data_sheets:
    print(f"No data sheets found in {src}", file=sys.stderr); sys.exit(1)

print(f"  flattening {len(data_sheets)} sheets: {', '.join(data_sheets)}", file=sys.stderr)

out_rows = 0
with open(dst, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f, quoting=csv.QUOTE_ALL)
    headers_written = False
    for sn in data_sheets:
        ws = wb[sn]
        sheet_iter = ws.iter_rows(values_only=True)
        next(sheet_iter)  # skip "Individual licensees" / "Organisation licensees" banner
        header = next(sheet_iter)
        if not headers_written:
            w.writerow(['SheetCategory'] + list(header))
            headers_written = True
        for row in sheet_iter:
            if row and any(v is not None for v in row):
                w.writerow([sn] + [str(v) if v is not None else '' for v in row])
                out_rows += 1

print(f"  wrote {out_rows} rows to {dst}", file=sys.stderr)
